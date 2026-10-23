"""Golden-value checks against the live database (skipped when unavailable).

BATTARAHALLI LAKE (112.08 latest / 134.57 mean / 151.04 worst) and IBBALUR
LAKE (115.03 latest / 143.94 worst) are the dashboard values over the default
last-12-months window (2020-12-28 .. 2021-12-28). Reports must reproduce them
exactly, so the exported workbook is compared to the API payload value by value.
"""
import asyncio
import io
import zipfile
from datetime import datetime
from pathlib import Path

import pytest
from dotenv import dotenv_values
from _pytest.outcomes import Skipped
from openpyxl import load_workbook

ENV = dotenv_values(Path(__file__).resolve().parents[1] / ".env")
pytestmark = pytest.mark.skipif(
    not ENV.get("MONGO_URI"), reason="live MongoDB connection not configured"
)

WINDOW = (datetime(2020, 12, 28), datetime(2021, 12, 28))
WATER_FILE = "waterKarna.csv"


def _database():
    from motor.motor_asyncio import AsyncIOMotorClient

    client = AsyncIOMotorClient(ENV["MONGO_URI"], serverSelectionTimeoutMS=8000)
    return client, client[ENV["DB_NAME"]]


def _water_dataset(db):
    return db.datasets.find_one({"filename": WATER_FILE})


async def _build(db, dataset, start=None, end=None):
    from app.reports.build import build_report_data

    user = {"_id": str(dataset.get("user_id")), "email": "reports@test"}
    return await build_report_data(
        db, user, dataset_ids=[str(dataset["_id"])], data_types=["water_quality"],
        start_date=start, end_date=end,
    )


def _run_build(start, end):
    async def run():
        client, db = _database()
        try:
            await db.command("ping")
            dataset = await _water_dataset(db)
            if not dataset:
                pytest.skip("waterKarna dataset not present")
            return await _build(db, dataset, start, end)
        finally:
            client.close()

    try:
        return asyncio.run(run())
    except Skipped:
        raise
    except Exception as exc:  # pragma: no cover - environment guard
        pytest.skip(f"live database unavailable: {exc}")


def _station(report, location):
    return next(row for row in report["stations"] if row["location"] == location)


def test_golden_values_match_dashboard_window():
    report = _run_build(*WINDOW)
    batt = _station(report, "BATTARAHALLI LAKE")
    ibb = _station(report, "IBBALUR LAKE")

    assert batt["latest_value"] == pytest.approx(112.08, abs=0.01)
    assert batt["mean_value"] == pytest.approx(134.57, abs=0.01)
    assert batt["max_value"] == pytest.approx(151.04, abs=0.01)
    assert batt["max_date"].date() == datetime(2021, 9, 14).date()

    assert ibb["latest_value"] == pytest.approx(115.03, abs=0.01)
    assert ibb["max_value"] == pytest.approx(143.94, abs=0.01)
    assert ibb["max_date"].date() == datetime(2021, 5, 19).date()

    assert report["meta"]["selection"]["station_count"] == 100
    assert report["meta"]["selection"]["row_count"] == 1114


def test_workbook_readback_equals_api_payload():
    from app.reports.excel import build_workbook

    report = _run_build(*WINDOW)
    workbook = load_workbook(io.BytesIO(build_workbook(report)))
    sheet = workbook["Station summary"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    header = [cell.value for cell in sheet[1]]

    def read(station):
        row = next(item for item in rows if item[2] == station)
        return dict(zip(header, row))

    for location, label in (("BATTARAHALLI LAKE", "Battarahalli Lake"),
                            ("IBBALUR LAKE", "Ibbalur Lake")):
        payload = _station(report, location)
        exported = read(label)
        assert exported["Latest"] == pytest.approx(payload["latest_value"])
        assert exported["Mean"] == pytest.approx(payload["mean_value"])
        assert exported["Min"] == pytest.approx(payload["min_value"])
        assert exported["Max"] == pytest.approx(payload["max_value"])
        assert exported["Scored samples"] == payload["scored_count"]
        assert exported["Max date"].date() == payload["max_date"].date()
        assert isinstance(exported["Mean"], float)
        assert isinstance(exported["Max date"], datetime)
        assert exported["Unit"] == "index"


def test_csv_row_count_matches_dataset_row_count():
    from app.reports.csv_zip import build_csv_zip

    async def run():
        client, db = _database()
        try:
            await db.command("ping")
            dataset = await _water_dataset(db)
            if not dataset:
                pytest.skip("waterKarna dataset not present")
            report = await _build(db, dataset)
            return await build_csv_zip(db, report)
        finally:
            client.close()

    try:
        data = asyncio.run(run())
    except Skipped:
        raise
    except Exception as exc:  # pragma: no cover - environment guard
        pytest.skip(f"live database unavailable: {exc}")

    archive = zipfile.ZipFile(io.BytesIO(data))
    name = next(item for item in archive.namelist() if item.startswith("data/"))
    with archive.open(name) as handle:
        total = sum(1 for _ in handle) - 1
    with archive.open("manifest.csv") as handle:
        lines = handle.read().decode("utf-8-sig").splitlines()
    assert int(lines[1].split(",")[5]) == total
    assert int(lines[1].split(",")[6]) == total == 3388


def test_date_range_matching_nothing_returns_empty_sections():
    report = _run_build(datetime(1900, 1, 1), datetime(1900, 1, 2))
    assert report["overview"]
    assert all(row["row_count"] == 0 for row in report["overview"])
    assert report["stations"] == []
    assert report["monthly"] == []
    assert report["anomalies"] == []
    assert report["meta"]["selection"]["row_count"] == 0

