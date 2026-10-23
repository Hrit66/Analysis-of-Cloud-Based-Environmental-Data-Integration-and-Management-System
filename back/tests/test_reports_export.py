"""Report renderers: workbook contract, CSV/ZIP streaming, scope and empty data."""
import asyncio
import csv
import io
import zipfile
from datetime import datetime
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook

from app.reports.build import ReportScopeError, build_report_data, resolve_datasets
from app.reports.csv_zip import build_csv_zip, stream_csv_zip
from app.reports.excel import build_workbook
from app.reports.pdf import build_pdf


def sample_report():
    return {
        "meta": {
            "title": "Environment Data Report",
            "generated_at": datetime(2026, 10, 20, 11, 30),
            "user_email": "imd.verifier@example.com",
            "user_name": "Verifier",
            "start_date": datetime(2020, 12, 28),
            "end_date": datetime(2021, 12, 28),
            "data_types": ["water_quality", "air_quality"],
            "dataset_ids": ["ds-water", "ds-air"],
            "datasets": [
                {
                    "id": "ds-water",
                    "name": "waterKarna.csv",
                    "filename": "waterKarna.csv",
                    "dataset_type": "water_quality",
                    "capabilities": ["water_quality"],
                    "tags": ["WQI computed"],
                    "row_count": 3388,
                    "created_at": datetime(2026, 10, 19),
                    "updated_at": None,
                },
                {
                    "id": "ds-air",
                    "name": "=city_day.csv",
                    "filename": "=city_day.csv",
                    "dataset_type": "air_quality",
                    "capabilities": ["air_quality"],
                    "tags": ["synthetic", "AQI computed"],
                    "row_count": 18265,
                    "created_at": datetime(2026, 10, 17),
                    "updated_at": None,
                },
            ],
            "selection": {
                "dataset_count": 2,
                "station_count": 2,
                "row_count": 21653,
                "anomaly_count": 3,
                "outlook_count": 1,
            },
            "notes": ["Water quality · waterKarna.csv: 1 stations, 12 rows."],
        },
        "overview": [
            {
                "dataset_id": "ds-water", "dataset_name": "waterKarna.csv",
                "data_type": "water_quality", "data_type_label": "Water quality",
                "station_count": 2, "row_count": 24,
                "start_date": datetime(2021, 1, 6), "end_date": datetime(2021, 12, 13),
                "coverage_pct": 100.0, "headline_parameter": "wqi",
                "headline_label": "WQI", "unit": "index",
                "missing": {"ph": 0.0, "do": 5.58},
                "not_in_source": ["tds"],
                "latest_value": 112.08, "latest_date": datetime(2021, 12, 13),
                "latest_station": "BATTARAHALLI LAKE", "category": "Very Poor",
                "mean_value": 124.8, "min_value": 47.21, "max_value": 164.01,
                "max_date": datetime(2021, 9, 14), "max_station": "BATTARAHALLI LAKE",
            },
            {
                "dataset_id": "ds-air", "dataset_name": "=city_day.csv",
                "data_type": "air_quality", "data_type_label": "Air quality",
                "station_count": 1, "row_count": 366,
                "start_date": datetime(2021, 1, 1), "end_date": datetime(2021, 12, 31),
                "coverage_pct": 100.0, "headline_parameter": "aqi",
                "headline_label": "AQI", "unit": "index",
                "missing": {"pm25": 0.0},
                "latest_value": 120.0, "latest_date": datetime(2021, 12, 31),
                "latest_station": "Bengaluru", "category": "Moderate",
                "mean_value": 110.0, "min_value": 40.0, "max_value": 320.0,
                "max_date": datetime(2021, 6, 1), "max_station": "Bengaluru",
            },
        ],
        "stations": [
            {
                "dataset_id": "ds-water", "dataset_name": "waterKarna.csv",
                "data_type": "water_quality", "data_type_label": "Water quality",
                "location": "BATTARAHALLI LAKE", "location_label": "Battarahalli Lake", "sample_count": 12,
                "first_date": datetime(2021, 1, 6, 8, 30), "last_date": datetime(2021, 12, 13, 8, 30),
                "latest_value": 112.08, "latest_category": "Very Poor",
                "mean_value": 134.57, "mean_category": "Unsuitable (>100)",
                "min_value": 107.69, "max_value": 151.04,
                "max_date": datetime(2021, 9, 14, 8, 30), "unit": "index",
                "rank": "80 of 98", "rank_position": 80, "rank_of": 98,
                "rank_samples": 12, "best_use_class": "D",
                "bacteriological_flag": "fails Class C (>5000 MPN/100 mL)",
            },
            {
                "dataset_id": "ds-air", "dataset_name": "=city_day.csv",
                "data_type": "air_quality", "data_type_label": "Air quality",
                "location": "=HYPERLINK(\"http://evil\")", "sample_count": 366,
                "first_date": datetime(2021, 1, 1), "last_date": datetime(2021, 12, 31),
                "latest_value": 120.0, "latest_category": "Moderate",
                "mean_value": 110.0, "mean_category": "Moderate",
                "min_value": 40.0, "max_value": 320.0,
                "max_date": datetime(2021, 6, 1), "unit": "index",
                "rank": None, "rank_position": None, "rank_of": None,
                "rank_samples": None, "best_use_class": None,
                "bacteriological_flag": None,
            },
        ],
        "monthly": [
            {
                "dataset_id": "ds-water", "dataset_name": "waterKarna.csv",
                "data_type": "water_quality", "parameter": "wqi",
                "parameter_label": "WQI", "unit": "index",
                "location": "BATTARAHALLI LAKE", "location_label": "Battarahalli Lake", "month": "2021-09",
                "month_date": datetime(2021, 9, 1),
                "mean": 151.04, "min": 151.04, "max": 151.04, "count": 1,
                "scored_count": 1, "partial": True,
            }
        ],
        "low_do": [
            {
                "dataset_id": "ds-water", "dataset_name": "waterKarna.csv",
                "data_type": "water_quality",
                "location": "GANGONDANA HALLI LAKE", "location_label": "Gangondana Halli Lake",
                "value": 0.4, "timestamp": datetime(2021, 8, 11, 8, 30),
            }
        ],
        "do_summary": [
            {
                "dataset_id": "ds-water", "dataset_name": "waterKarna.csv",
                "critical": 26, "low": 39, "below_one": 2,
            }
        ],
        "anomalies": [
            {
                "dataset_id": "ds-air", "dataset_name": "=city_day.csv",
                "data_type": "air_quality", "data_type_label": "Air quality",
                "location": "Bengaluru", "timestamp": datetime(2021, 6, 1),
                "parameter": "pm25", "parameter_label": "PM2.5", "unit": "µg/m³",
                "value": 320.0, "expected_low": 10.0, "expected_high": 90.0,
                "baseline": 45.0, "severity": "extreme", "kind": None,
                "detector": "isolation_forest", "score": -0.31,
            },
            {
                "dataset_id": "ds-water", "dataset_name": "waterKarna.csv",
                "data_type": "water_quality", "data_type_label": "Water quality",
                "location": "BATTARAHALLI LAKE", "timestamp": datetime(2021, 9, 14, 8, 30),
                "parameter": "wqi", "parameter_label": "WQI", "unit": "index",
                "value": 151.04, "expected_low": 20.0, "expected_high": 140.0,
                "baseline": 80.0, "severity": "high", "kind": "unusually_high",
                "detector": "robust_zscore", "score": 4.2,
            },
        ],
        "anomaly_totals": [
            {
                "dataset_id": "ds-water", "dataset_name": "waterKarna.csv",
                "data_types": ["water_quality"], "data_types_label": "Water quality",
                "flagged": 1, "scanned": 24, "rate_pct": 4.17, "parameters": [],
            },
            {
                "dataset_id": "ds-air", "dataset_name": "=city_day.csv",
                "data_types": ["air_quality"], "data_types_label": "Air quality",
                "flagged": 2, "scanned": 366, "rate_pct": 0.55, "parameters": [],
            },
        ],
        "outlooks": [
            {
                "dataset_id": "ds-water", "dataset_name": "waterKarna.csv",
                "data_type": "water_quality", "data_type_label": "Water quality",
                "location": "BATTARAHALLI LAKE", "parameter": "wqi",
                "parameter_label": "WQI", "unit": "index",
                "method": "SimpleExponentialSmoothing",
                "target_date": datetime(2022, 1, 1),
                "forecast_value": 115.0, "lower_bound": 100.0, "upper_bound": 130.0,
                "horizon_steps": 3, "horizon_label": 3, "mae": 0.98,
                "baseline_mae": 1.356, "skill_pct": 27.7, "sample_count": 39,
                "note": "Monthly outlook from 39 points.",
            }
        ],
        "outlook_coverage": [
            {
                "dataset_id": "ds-water", "dataset_name": "waterKarna.csv",
                "stations_with_outlook": 1, "station_count": 2, "sparse_slots": 1,
                "unexplained_slots": 0, "min_months": 24,
                "missing_stations": ["Ibbalur Lake"],
            }
        ],
        "data_quality": [
            {
                "dataset_id": "ds-water", "dataset_name": "waterKarna.csv",
                "data_type": "water_quality",
                "rows_read": 3434, "rows_stored": 3388, "rows_dropped": 0,
                "duplicates_removed": 46, "columns_dropped": 78,
                "source_format": "tabular",
                "notes": ["46 duplicate station-dates averaged."],
            }
        ],
        "notes": ["waterKarna.csv · Water quality: 100 stations, 1,111 rows."],
    }


class FakeCursor:
    def __init__(self, rows):
        self._rows = list(rows)

    async def to_list(self, length=None):
        return self._rows if length is None else self._rows[:length]


def _matches(doc, query):
    for key, value in query.items():
        if key == "$or":
            if not any(_matches(doc, clause) for clause in value):
                return False
        elif key == "$and":
            if not all(_matches(doc, clause) for clause in value):
                return False
        elif isinstance(value, dict) and "$in" in value:
            if doc.get(key) not in value["$in"]:
                return False
        elif isinstance(value, dict) and "$exists" in value:
            if (key in doc) != value["$exists"]:
                return False
        elif doc.get(key) != value:
            return False
    return True


class FakeDatasets:
    def __init__(self, docs):
        self.docs = docs

    def find(self, query):
        return FakeCursor([doc for doc in self.docs if _matches(doc, query)])


class FakeDatabase:
    def __init__(self, docs=None):
        self.datasets = FakeDatasets(docs or [])


def test_empty_selection_builds_an_empty_report_without_queries():
    report = asyncio.run(build_report_data(object(), {"_id": "u1"}, dataset_ids=[]))
    assert report["meta"]["selection"]["dataset_count"] == 0
    assert report["overview"] == [] and report["stations"] == []
    assert report["anomalies"] == [] and report["outlooks"] == []


def test_foreign_dataset_id_is_refused():
    db = FakeDatabase([{
        "_id": "000000000000000000000001", "user_id": "someone-else",
        "status": "completed", "capabilities": ["water_quality"],
    }])
    with pytest.raises(ReportScopeError):
        asyncio.run(resolve_datasets(db, "u1", ["000000000000000000000001"], None))


def test_authenticated_listing_only_returns_owned_datasets():
    own = {"_id": "a", "user_id": "u1", "status": "completed",
           "capabilities": ["water_quality"], "dataset_type": "water_quality",
           "filename": "mine.csv"}
    other = {"_id": "b", "user_id": "u2", "status": "completed",
             "capabilities": ["air_quality"], "dataset_type": "air_quality",
             "filename": "theirs.csv"}
    selected = asyncio.run(resolve_datasets(FakeDatabase([own, other]), "u1", None, None))
    assert [str(doc["_id"]) for doc, _ in selected] == ["a"]


def test_workbook_contract_and_native_types():
    data = build_workbook(sample_report())
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == [
        "About", "Overview", "Station summary", "Monthly", "Anomalies",
        "Outlooks", "Data quality",
    ]
    assert all(len(name) <= 31 for name in wb.sheetnames)

    ws = wb["Station summary"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    batt = next(row for row in ws.iter_rows(min_row=2, values_only=True)
                if row[2] == "Battarahalli Lake")
    assert batt[6] == pytest.approx(112.08)
    assert batt[7] == "Very Poor"
    assert batt[8] == pytest.approx(134.57)
    assert batt[9] == "Unsuitable (>100)"
    assert batt[11] == pytest.approx(151.04)
    assert isinstance(batt[4], datetime)
    assert batt[14] == 80
    assert batt[15] == 98

    monthly = wb["Monthly"]
    month_header = [cell.value for cell in monthly[1]]
    assert "Scored" in month_header
    assert "Partial month" in month_header
    month_row = next(item for item in monthly.iter_rows(min_row=2, values_only=True) if item[3] == datetime(2021, 9, 1))
    assert month_row[month_header.index("Scored")] == 1
    assert month_row[month_header.index("Partial month")] == "yes"

    outlooks = wb["Outlooks"]
    outlook_header_row = next(
        index for index, row in enumerate(outlooks.iter_rows(values_only=True), start=1)
        if row[0] == "Dataset"
    )
    assert outlook_header_row == 1
    assert outlooks.cell(outlooks.max_row, 1).value.startswith("waterKarna.csv:")
    assert outlooks.auto_filter.ref.split(":")[0] == f"A{outlook_header_row}"

    about = wb["About"]
    texts = [cell.value for row in about.iter_rows() for cell in row if isinstance(cell.value, str)]
    assert any(value.startswith("'=city_day.csv") for value in texts)
    assert any("28 Dec 2020 – 28 Dec 2021" in value for value in texts)
    assert not any(value.strip().lower() == "nan" for value in texts)


def test_workbook_sanitises_formula_like_station_names():
    data = build_workbook(sample_report())
    wb = load_workbook(io.BytesIO(data))
    stations = [row[2] for row in wb["Station summary"].iter_rows(min_row=2, values_only=True)]
    assert any(value.startswith("'=HYPERLINK") for value in stations)


class FakeRecordCursor:
    def __init__(self, docs):
        self.docs = iter(docs)

    def sort(self, *args, **kwargs):
        return self

    def __aiter__(self):
        return self

    async def __anext__(self):
        try:
            return next(self.docs)
        except StopIteration:
            raise StopAsyncIteration


class FakeRecords:
    def __init__(self, docs):
        self.docs = docs

    def find(self, query):
        return FakeRecordCursor(self.docs)


class FakeCSVDatabase:
    def __init__(self, docs):
        self.water_quality = FakeRecords(docs)


def test_csv_zip_streams_one_file_per_dataset_with_manifest():
    docs = [
        {"timestamp": datetime(2021, 1, 6, 8, 30), "location": "BATTARAHALLI LAKE",
         "ph": 7.1, "do": 5.2, "wqi": 112.08, "wqi_category": "Poor"},
        {"timestamp": datetime(2021, 2, 6, 8, 30), "location": "BATTARAHALLI LAKE",
         "ph": None, "do": None, "wqi": None, "wqi_category": None},
        {"timestamp": datetime(2021, 3, 6, 8, 30), "location": "=cmd()",
         "ph": 8.0, "do": 6.0, "wqi": 90.0, "wqi_category": "Poor"},
    ]
    report = sample_report()
    report["meta"]["datasets"] = [report["meta"]["datasets"][0]]
    report["meta"]["datasets"][0]["row_count"] = 3
    database = FakeCSVDatabase(docs)

    async def collect():
        return [chunk async for chunk in stream_csv_zip(database, report)]

    chunks = asyncio.run(collect())
    archive = zipfile.ZipFile(io.BytesIO(b"".join(chunks)))
    names = archive.namelist()
    assert "manifest.csv" in names
    data_names = [name for name in names if name.startswith("data/")]
    assert len(data_names) == 1

    raw = archive.open(data_names[0]).read()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert len(rows) == 4
    header = rows[0]
    assert any("WQI" in column for column in header)
    assert "Timestamp (UTC)" in header
    batt = next(row for row in rows[1:] if row[1] == "BATTARAHALLI LAKE")
    assert batt[0] == "2021-01-06T08:30:00"
    assert row_value(batt, header, "pH") == "7.1"
    assert row_value(rows[2], header, "pH") == ""
    assert next(row for row in rows[1:] if "cmd" in row[1])[1].startswith("'=")

    manifest = list(csv.reader(io.StringIO(archive.open("manifest.csv").read().decode("utf-8-sig"))))
    assert manifest[1][5] == "3"
    assert manifest[1][6] == "3"


def row_value(row, header, label):
    index = next(i for i, column in enumerate(header) if column == label or column.startswith(f"{label} "))
    return row[index]


def test_pdf_contains_findings_and_sections():
    data = build_pdf(sample_report())
    assert data.startswith(b"%PDF")
    try:
        import subprocess
        text = subprocess.run(
            ["pdftotext", "-", "-"], input=data, capture_output=True, check=True,
        ).stdout.decode("utf-8", "replace")
    except (FileNotFoundError, subprocess.CalledProcessError):
        return
    for phrase in ("Environment Data Report", "Key findings", "Battarahalli Lake",
                   "Not in source", "Best 5 by mean",
                   "Lowest dissolved-oxygen samples in the window",
                   "the largest anomaly score is",
                   "Methodology and limitations", "Data quality appendix"):
        assert phrase in text
    assert "µg/m³" in text


def test_preview_and_export_require_authentication():
    from fastapi.testclient import TestClient
    from app.main import app

    client = TestClient(app)
    assert client.post("/api/reports/preview", json={"dataset_ids": "all"}).status_code == 401
    assert client.post(
        "/api/reports/export", json={"dataset_ids": "all", "format": "xlsx"}
    ).status_code == 401


def test_export_rejects_foreign_dataset():
    from fastapi.testclient import TestClient
    from app.api.auth import get_current_user_dep, get_database
    from app.main import app

    foreign_id = "000000000000000000000001"
    fake = FakeDatabase([{
        "_id": foreign_id, "user_id": "someone-else", "status": "completed",
        "capabilities": ["water_quality"], "dataset_type": "water_quality",
        "filename": "theirs.csv", "row_count": 10,
    }])
    app.dependency_overrides[get_database] = lambda: fake
    app.dependency_overrides[get_current_user_dep] = lambda: {
        "_id": "u1", "email": "u1@example.com",
    }
    try:
        client = TestClient(app)
        response = client.post(
            "/api/reports/export",
            json={"dataset_ids": [foreign_id], "format": "xlsx"},
        )
        preview = client.post(
            "/api/reports/preview",
            json={"dataset_ids": [foreign_id]},
        )
    finally:
        app.dependency_overrides.clear()
    assert response.status_code == 404
    assert preview.status_code == 404
    assert "not available" in response.json()["detail"]


def test_inclusive_end_expands_midnight_only():
    from datetime import time as clock_time
    from app.services.dashboard_queries import inclusive_end

    assert inclusive_end(None) is None
    expanded = inclusive_end(datetime(2021, 12, 28))
    assert expanded.time() == clock_time.max
    exact = datetime(2021, 12, 28, 8, 30)
    assert inclusive_end(exact) == exact


def test_display_location_cleans_typos_and_duplicate_suffixes():
    from app.reports.build import display_location

    assert display_location("JIGANI TNAK") == "Jigani Tank"
    assert display_location("NETRAVATHI AT U/S OF UPIINGADY TOWN") == (
        "Netravathi at U/S of Uppinagady Town"
    )
    assert display_location("ANDHRAHALLI LAKE ANDHRAHALLI BENGALURU") == "Andhrahalli Lake"
    assert display_location("CHIKKABANVARA LAKE CHIKKABANWARA BENGALURU NORTH") == (
        "Chikkabanvara Lake"
    )


def test_low_do_flags_use_absolute_thresholds():
    from app.services.dashboard_queries import _low_do_flag

    assert _low_do_flag(None) is None
    assert _low_do_flag(1.9) == "critical low DO (<2 mg/L)"
    assert _low_do_flag(3.9) == "low DO (<4 mg/L)"
    assert _low_do_flag(4.0) is None


def test_data_quality_notes_count_what_they_claim():
    from app.reports.build import _data_quality_row

    row = _data_quality_row({
        "_id": "ds-water",
        "filename": "waterKarna.csv",
        "dataset_type": "water_quality",
        "row_count": 100,
        "cleaning_report": {
            "rows_read": 120,
            "missing_markers_found": {"ph": 2, "do": 3},
            "columns_dropped": ["a", "b"],
            "columns_renamed": {"x": "ph"},
            "synonym_columns_merged": {"do": ["_do", "d_o"]},
        },
    })
    notes = " ".join(row.get("notes") or [])
    assert "5 blank/NA missing-value markers were blanked across 2 stored parameters" in notes
    assert "2 unused/empty source columns dropped; 1 renamed and 2 synonym aliases merged" in notes
    assert row["columns_dropped"] == 2
    assert row["columns_renamed"] == 1
    assert row["synonym_columns_merged"] == 2
