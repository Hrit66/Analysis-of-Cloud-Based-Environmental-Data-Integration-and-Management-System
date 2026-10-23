"""Regenerate the dirty-data fixtures used by tests/test_pipeline_dirty.py.

Run from back/:
    ./venv/bin/python tests/fixtures/generate_fixtures.py
"""
from pathlib import Path

from openpyxl import Workbook

FIXTURES = Path(__file__).resolve().parent


def write_text(name: str, text: str, encoding: str = "utf-8") -> None:
    (FIXTURES / name).write_bytes(text.encode(encoding))


def csv(rows) -> str:
    return "\n".join(",".join(str(cell) for cell in row) for row in rows) + "\n"


def main() -> None:
    write_text(
        "column_variants.csv",
        csv(
            [
                ["Date", "city", "PM2.5 (µg/m³)", "pm_10", "NO2", "SO2", "CO", "O3", "Temperature (°C)", "Precipitation (mm)"],
                ["2026-10-01", "Bengaluru", "45", "80", "30", "10", "1.2", "40", "28.5", "0"],
                ["2026-10-02", "Bengaluru", "52", "91", "33", "11", "1.4", "44", "29.1", "2.5"],
            ]
        ),
    )

    write_text(
        "missing_markers.csv",
        csv(
            [
                ["timestamp", "location", "pm25", "pm10", "no2", "so2", "co", "o3"],
                ["2026-10-01 00:00:00", "MarkerCity", "NA", "80", "30", "10", "1.2", "40"],
                ["2026-10-01 01:00:00", "MarkerCity", "-", "80", "30", "10", "1.2", "40"],
                ["2026-10-01 02:00:00", "MarkerCity", "None", "80", "30", "10", "1.2", "40"],
                ["2026-10-01 03:00:00", "MarkerCity", "null", "80", "30", "10", "1.2", "40"],
                ["2026-10-01 04:00:00", "MarkerCity", "", "80", "30", "10", "1.2", "40"],
                ["2026-10-01 05:00:00", "MarkerCity", "--", "80", "30", "10", "1.2", "40"],
                ["2026-10-01 06:00:00", "MarkerCity", "N/A", "80", "30", "10", "1.2", "40"],
            ]
        ),
    )

    write_text(
        "sentinels_impossible.csv",
        csv(
            [
                ["timestamp", "location", "pm25", "pm10", "temperature", "precipitation"],
                ["2026-10-01 00:00:00", "SentinelCity", "-999", "9999", "25", "5"],
                ["2026-10-02 00:00:00", "SentinelCity", "-5", "2000", "500", "-1"],
            ]
        ),
    )

    write_text(
        "below_detection.csv",
        csv(
            [
                ["timestamp", "location", "pm25", "pm10", "no2"],
                ["2026-10-01 00:00:00", "BDLCity", "<2", "BDL", "30"],
                ["2026-10-01 01:00:00", "BDLCity", "ND", "<1.5", "32"],
            ]
        ),
    )

    write_text(
        "impossible_water.csv",
        csv(
            [
                ["timestamp", "location", "pH", "turbidity", "TDS", "DO"],
                ["2026-10-01 00:00:00", "RiverPoint", "7.2", "5", "300", "6.5"],
                ["2026-10-01 01:00:00", "RiverPoint", "20", "-50", "99999", "200"],
            ]
        ),
    )

    write_text(
        "date_formats.csv",
        csv(
            [
                ["timestamp", "location", "pm25"],
                ["2026-10-11", "DateA", "10"],
                ["11-10-2026", "DateB", "11"],
                ["10/11/2026", "DateC", "12"],
                ["11 Oct 2026 14:00", "DateD", "13"],
            ]
        ),
    )

    write_text(
        "timezones.csv",
        csv(
            [
                ["timestamp", "location", "pm25"],
                ["2026-10-11T10:00:00Z", "TZ-Z", "10"],
                ["2026-10-11 15:30 IST", "TZ-IST", "11"],
                ["2026-10-11 10:00 UTC", "TZ-UTC", "12"],
                ["2026-10-11 10:00:00+05:30", "TZ-OFFSET", "13"],
                ["2026-10-11 10:00:00", "TZ-NAIVE", "14"],
            ]
        ),
    )

    write_text(
        "duplicates.csv",
        csv(
            [
                ["timestamp", "location", "pm25"],
                ["2026-10-11T00:00:00Z", "DupCity", "10"],
                ["2026-10-11T00:00:00Z", "DupCity", "10"],
                ["2026-10-11T00:00:00Z", "DupCity", "12"],
                ["2026-10-11T01:00:00Z", "DupCity", "20"],
            ]
        ),
    )

    write_text(
        "wrong_type_weather.csv",
        csv(
            [
                ["date", "city", "temperature", "precipitation"],
                ["2026-10-01", "Bengaluru", "28.5", "0"],
                ["2026-10-02", "Bengaluru", "29.1", "2.5"],
            ]
        ),
    )

    workbook = Workbook()
    notes = workbook.active
    notes.title = "Notes"
    notes["A1"] = "Field notes — not data"
    notes["B2"] = "ignore me"

    data = workbook.create_sheet("Data")
    data.merge_cells("A1:C1")
    data["A1"] = "Air Quality Report 2026"
    data["A3"], data["B3"], data["C3"] = "Date", "Station", "PM2.5"
    data["A4"], data["B4"], data["C4"] = "2026-10-01", "ExcelCity", 42
    data["A5"], data["B5"], data["C5"] = "2026-10-02", "ExcelCity", 47
    xlsx_bytes = None
    workbook.save(FIXTURES / "excel_quirks.xlsx")
    xlsx_bytes = (FIXTURES / "excel_quirks.xlsx").read_bytes()
    (FIXTURES / "renamed_excel.csv").write_bytes(xlsx_bytes)

    write_text(
        "json_records.json",
        """[
  {"timestamp": "2026-10-01 00:00:00", "location": "JSONCity", "pm25": 18, "pm10": 30},
  {"timestamp": "2026-10-01 01:00:00", "location": "JSONCity", "pm25": 21, "pm10": 34}
]
""",
    )

    write_text(
        "json_nested.json",
        """[
  {"timestamp": "2026-10-01 00:00:00", "location": "NestedCity", "readings": {"pm25": 12, "pm10": 20, "no2": 7}},
  {"timestamp": "2026-10-01 01:00:00", "location": "NestedCity", "readings": {"pm25": 15, "pm10": 22, "no2": 9}}
]
""",
    )

    write_text(
        "json_data_key.json",
        """{
  "status": "ok",
  "data": [
    {"timestamp": "2026-10-01 00:00:00", "location": "WrappedCity", "pm25": 33, "pm10": 55}
  ]
}
""",
    )

    write_text(
        "long_format.csv",
        csv(
            [
                ["location", "datetime", "parameter", "value", "units"],
                ["OpenAQStation", "2026-10-11T00:00:00Z", "pm25", "12", "ug/m3"],
                ["OpenAQStation", "2026-10-11T00:00:00Z", "pm10", "24", "ug/m3"],
                ["OpenAQStation", "2026-10-11T00:00:00Z", "no2", "8", "ug/m3"],
                ["OpenAQStation", "2026-10-11T00:00:00Z", "o3", "40", "ug/m3"],
                ["OpenAQStation", "2026-10-11T00:00:00Z", "pressure", "1013", "hPa"],
                ["OpenAQStation", "2026-10-11T01:00:00Z", "pm25", "13", "ug/m3"],
                ["OpenAQStation", "2026-10-11T01:00:00Z", "pm10", "26", "ug/m3"],
                ["OpenAQStation", "2026-10-11T01:00:00Z", "no2", "9", "ug/m3"],
                ["OpenAQStation", "2026-10-11T01:00:00Z", "o3", "42", "ug/m3"],
            ]
        ),
    )

    write_text(
        "bom_utf8.csv",
        csv(
            [
                ["timestamp", "location", "PM2.5 (µg/m³)"],
                ["2026-10-01 00:00:00", "BomCity", "25"],
            ]
        ),
    )
    bom_path = FIXTURES / "bom_utf8.csv"
    bom_path.write_bytes(b"\xef\xbb\xbf" + bom_path.read_bytes())

    write_text(
        "mojibake_utf8.csv",
        csv(
            [
                ["timestamp", "location", "PM2.5 (Âµg/mÂ³)"],
                ["2026-10-01 00:00:00", "MojibakeCity", "30"],
            ]
        ),
    )

    write_text(
        "latin1_encoded.csv",
        csv(
            [
                ["timestamp", "location", "PM2.5 (µg/m³)", "Temp (°C)"],
                ["2026-10-01 00:00:00", "Latin1City", "25", "28"],
            ]
        ),
        encoding="latin-1",
    )

    (FIXTURES / "empty.csv").write_bytes(b"")

    write_text("header_only.csv", csv([["timestamp", "location", "pm25"]]))

    write_text(
        "single_row.csv",
        csv(
            [
                ["timestamp", "location", "pm25"],
                ["2026-10-01 00:00:00", "SoloCity", "17"],
            ]
        ),
    )

    write_text(
        "actually_json.txt",
        """[
  {"timestamp": "2026-10-01 00:00:00", "location": "SniffCity", "pm25": 5, "pm10": 9}
]
""",
    )

    write_text(
        "imd_rain_testcity.csv",
        csv(
            [
                ["DateTime", "12.97 77.59"],
                ["2023-01-01", "0.0000"],
                ["2023-01-02", "12.5000"],
                ["2023-01-03", "-999"],
                ["2023-01-04", "-2.5"],
            ]
        ),
    )

    write_text(
        "imd_tmax_testcity.csv",
        csv(
            [
                ["DateTime", "12.97 77.59"],
                ["2023-01-01", "29.0"],
                ["2023-01-02", "30.0"],
                ["2023-01-03", "31.0"],
            ]
        ),
    )

    write_text(
        "imd_tmin_testcity.csv",
        csv(
            [
                ["DateTime", "12.97 77.59"],
                ["2023-01-01", "15.0"],
                ["2023-01-02", "16.0"],
                ["2023-01-03", "17.0"],
                ["2023-01-04", "18.0"],
            ]
        ),
    )

    write_text(
        "imd_rain_corrupt_bengaluru.csv",
        csv(
            [
                ["DateTime", "12.97 77.59"],
                ["2023-01-01", "0.0"],
                ["2023-01-02", "-999"],
                ["2023-01-03", "-2.5"],
                ["2023-01-04", "10.2"],
            ]
        ),
    )

    write_text(
        "imd_broken_missing_column.csv",
        csv([["DateTime"], ["2023-01-01"], ["2023-01-02"]]),
    )

    write_text(
        "imd_station_data.csv",
        csv([["DateTime", "12.97 77.59"], ["2023-01-01", "5.0"]]),
    )

    print(f"fixtures written to {FIXTURES}")


if __name__ == "__main__":
    main()
